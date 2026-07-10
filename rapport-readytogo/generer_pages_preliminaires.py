#!/usr/bin/env python3
"""Génère les pages préliminaires du rapport ReadyToGo (Word académique)."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pathlib import Path


OUTPUT = Path(__file__).resolve().parent / "ReadyToGo_Rapport_Management_Projet.docx"


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_format(
    paragraph,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    space_before=0,
    space_after=6,
    line_spacing=1.5,
    first_line_indent=None,
):
    pf = paragraph.paragraph_format
    pf.alignment = alignment
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)


def add_page_number(paragraph):
    """Insère un champ PAGE dans le paragraphe (pied de page)."""
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_run_font(run, size=11)


def add_toc_field(paragraph):
    """Insère un champ TOC automatique (à actualiser dans Word)."""
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = ' TOC \\o "1-3" \\h \\z \\u '

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)

    placeholder = paragraph.add_run(
        "[Sommaire automatique — clic droit > Mettre à jour les champs dans Microsoft Word]"
    )
    set_run_font(placeholder, size=11, italic=True, color=RGBColor(0x66, 0x66, 0x66))

    run2 = paragraph.add_run()
    run2._r.append(fld_char_end)


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_section_title(doc, text, level=1):
    """Titre de section préliminaire (style Heading pour le TOC)."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    set_paragraph_format(heading, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=12)
    return heading


def add_body_paragraph(doc, text, first_indent=0.75):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=first_indent, space_after=10)
    run = p.add_run(text)
    set_run_font(run, size=12)
    return p


def add_centered_line(doc, text, size=12, bold=False, italic=False, space_before=0, space_after=6, color=None):
    p = doc.add_paragraph()
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=space_before, space_after=space_after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_page_break(doc):
    doc.add_page_break()


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for i, size in [(1, 16), (2, 14), (3, 12)]:
        style = styles[f"Heading {i}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        style.paragraph_format.line_spacing = 1.5


def setup_section(section, different_first_page=True):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.different_first_page_header_footer = different_first_page

    # Pied de page (pages suivantes) : numéro centré
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    set_paragraph_format(fp, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.0)
    add_page_number(fp)

    # Première page (garde) : pas de numéro
    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    ffp = first_footer.paragraphs[0]
    ffp.clear()


def build_cover_page(doc):
    # Espace haut
    for _ in range(2):
        p = doc.add_paragraph()
        set_paragraph_format(p, space_after=0, line_spacing=1.0)

    add_centered_line(
        doc,
        "Université Abdelmalek Essaâdi",
        size=14,
        bold=True,
        space_after=4,
        color=RGBColor(0x1F, 0x4E, 0x79),
    )
    add_centered_line(
        doc,
        "Faculté des Sciences Juridiques, Économiques et Sociales de Tanger",
        size=12,
        bold=False,
        space_after=4,
    )
    add_centered_line(doc, "Master : IA pour l’Économie Numérique et la Gestion", size=12, space_after=4)
    add_centered_line(doc, "Module : Management de Projet", size=12, space_after=18)

    line = doc.add_paragraph()
    set_paragraph_format(line, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=18)
    add_horizontal_line(line)

    add_centered_line(doc, "Rapport de Management de Projet", size=16, bold=True, space_before=12, space_after=18)

    add_centered_line(
        doc,
        "ReadyToGo",
        size=28,
        bold=True,
        space_before=6,
        space_after=14,
        color=RGBColor(0x1F, 0x4E, 0x79),
    )

    add_centered_line(
        doc,
        "Service intelligent de trottinettes électriques pour améliorer\n"
        "la mobilité urbaine à Tanger dans la perspective\n"
        "de la Coupe du Monde 2030",
        size=13,
        italic=True,
        space_before=6,
        space_after=24,
    )

    line2 = doc.add_paragraph()
    set_paragraph_format(line2, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=18)
    add_horizontal_line(line2)

    # Bloc auteurs / encadrant
    for _ in range(1):
        p = doc.add_paragraph()
        set_paragraph_format(p, space_after=0, line_spacing=1.0)

    add_centered_line(doc, "Réalisé par :", size=12, bold=True, space_before=12, space_after=4)
    for name in ["Nom 1", "Nom 2", "Nom 3", "Nom 4"]:
        add_centered_line(doc, name, size=12, space_after=2)

    add_centered_line(doc, "Encadré par :", size=12, bold=True, space_before=18, space_after=4)
    add_centered_line(doc, "Pr. EL HAJJAJI", size=12, space_after=24)

    add_centered_line(doc, "Année universitaire :", size=12, bold=True, space_before=12, space_after=4)
    add_centered_line(doc, "2025 / 2026", size=12, space_after=6)

    # Emplacement logo facultatif
    note = doc.add_paragraph()
    set_paragraph_format(note, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=0)
    run = note.add_run(
        "[Emplacement suggéré : logo de l’Université / de la Faculté en en-tête de la page de garde]"
    )
    set_run_font(run, size=9, italic=True, color=RGBColor(0x88, 0x88, 0x88))


def build_remerciements(doc):
    add_page_break(doc)
    add_section_title(doc, "Remerciements")

    paragraphs = [
        (
            "Nous tenons à exprimer nos sincères remerciements à notre professeur encadrant, "
            "Pr. EL HAJJAJI, pour la qualité de son accompagnement, ses orientations méthodologiques "
            "et les connaissances transmises dans le cadre du module de Management de Projet."
        ),
        (
            "Nous remercions également l’ensemble des membres de notre équipe pour leur implication, "
            "leur esprit de collaboration et leur contribution à la construction progressive de ce projet. "
            "Ce travail nous a permis de mettre en pratique les outils fondamentaux du management de projet, "
            "notamment le cadrage, la planification, la gestion des risques, le suivi des indicateurs "
            "et la préparation des livrables."
        ),
        (
            "Enfin, nous remercions toutes les personnes qui ont contribué, directement ou indirectement, "
            "à enrichir notre réflexion autour des problématiques de mobilité urbaine, de développement "
            "durable et d’amélioration des services de transport à Tanger."
        ),
    ]
    for text in paragraphs:
        add_body_paragraph(doc, text)


def build_resume(doc):
    add_page_break(doc)
    add_section_title(doc, "Résumé exécutif")

    paragraphs = [
        (
            "Le présent rapport porte sur le projet ReadyToGo, un service intelligent de trottinettes "
            "électriques en libre-service destiné à améliorer la mobilité urbaine à Tanger. Le projet "
            "s’inscrit dans un contexte marqué par la croissance de la circulation urbaine, la congestion "
            "des axes principaux et les besoins futurs liés à l’accueil d’événements internationaux, "
            "notamment la Coupe du Monde 2030."
        ),
        (
            "ReadyToGo propose une solution de mobilité légère, flexible et accessible, permettant aux "
            "habitants, étudiants, touristes et visiteurs de se déplacer rapidement sur des trajets courts "
            "et moyens. Le service repose sur une application mobile, des trottinettes électriques "
            "géolocalisées, des zones de stationnement autorisées et un système de paiement simple."
        ),
        (
            "L’objectif du projet est de concevoir, tester et déployer progressivement un service de "
            "micromobilité à Tanger, en commençant par une phase pilote limitée à certaines zones "
            "stratégiques telles que la corniche, le centre-ville, la gare Tanger-Ville et les quartiers "
            "touristiques. À moyen terme, le service pourra être étendu vers l’aéroport Ibn Battouta, "
            "les zones hôtelières et le Grand Stade de Tanger."
        ),
        (
            "Ce rapport présente le projet selon une démarche complète de management de projet. Il traite "
            "le cadrage, les objectifs SMART, les parties prenantes, le périmètre, le cahier des charges, "
            "la méthode de gestion adoptée, la planification, le budget prévisionnel, la gestion des risques, "
            "les indicateurs de suivi, la communication, les scénarios d’usage et la clôture du projet."
        ),
        (
            "La méthodologie retenue est une approche hybride combinant la rigueur de la gestion classique "
            "pour le cadrage, le budget et les autorisations, avec l’agilité pour le développement de "
            "l’application mobile, les tests utilisateurs et l’amélioration continue du service."
        ),
    ]
    for text in paragraphs:
        add_body_paragraph(doc, text)


def build_abstract(doc):
    add_page_break(doc)
    add_section_title(doc, "Abstract")

    paragraphs = [
        (
            "This report presents ReadyToGo, a smart electric scooter-sharing service designed to improve "
            "urban mobility in Tangier, Morocco. The project responds to increasing traffic congestion, "
            "growing mobility needs, and the city’s preparation for major international events such as "
            "the 2030 FIFA World Cup."
        ),
        (
            "ReadyToGo aims to provide a flexible, affordable and environmentally friendly mobility solution "
            "for residents, students, tourists and visitors. The service is based on a mobile application, "
            "GPS-enabled electric scooters, designated parking zones, digital payment and operational "
            "maintenance teams."
        ),
        (
            "The project will be implemented progressively, starting with a pilot phase in strategic areas "
            "such as the city center, the corniche, Tangier-Ville railway station and tourist districts. "
            "Later phases may extend the service to Ibn Battouta Airport, hotel areas and the Grand Stadium "
            "of Tangier."
        ),
        (
            "This report applies project management concepts including project framing, SMART objectives, "
            "stakeholder analysis, scope definition, specifications, hybrid project methodology, WBS, "
            "roadmap, Gantt planning, PERT analysis, budgeting, risk management, KPI monitoring, "
            "communication planning and project closure."
        ),
        (
            "The selected methodology is hybrid: a traditional approach is used for feasibility, "
            "authorizations, budgeting and governance, while an agile approach is used for the mobile "
            "application, user testing and continuous improvement."
        ),
    ]
    for text in paragraphs:
        add_body_paragraph(doc, text)


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "1F4E79")
        tcBorders.append(element)
    tcPr.append(tcBorders)


def build_sigles(doc):
    add_page_break(doc)
    add_section_title(doc, "Liste des sigles et abréviations")

    intro = doc.add_paragraph()
    set_paragraph_format(intro, space_after=12, first_line_indent=0)
    run = intro.add_run(
        "Les sigles et abréviations utilisés dans le présent rapport sont présentés ci-dessous."
    )
    set_run_font(run, size=12)

    data = [
        ("BRT", "Bus Rapid Transit"),
        ("CNDP", "Commission Nationale de Contrôle de la Protection des Données à Caractère Personnel"),
        ("GPS", "Global Positioning System"),
        ("IA", "Intelligence Artificielle"),
        ("KPI", "Key Performance Indicator"),
        ("MVP", "Minimum Viable Product"),
        ("PERT", "Program Evaluation and Review Technique"),
        ("RACI", "Responsible, Accountable, Consulted, Informed"),
        ("SMART", "Spécifique, Mesurable, Atteignable, Réaliste, Temporel"),
        ("WBS", "Work Breakdown Structure"),
    ]

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0].cells
    hdr[0].text = ""
    hdr[1].text = ""
    for i, label in enumerate(["Sigle", "Signification"]):
        hdr[i].paragraphs[0].clear()
        set_paragraph_format(hdr[i].paragraphs[0], alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, line_spacing=1.15)
        run = hdr[i].paragraphs[0].add_run(label)
        set_run_font(run, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(hdr[i], "1F4E79")
        set_cell_borders(hdr[i])

    for sigle, signification in data:
        row = table.add_row().cells
        for cell, value, align in [
            (row[0], sigle, WD_ALIGN_PARAGRAPH.CENTER),
            (row[1], signification, WD_ALIGN_PARAGRAPH.LEFT),
        ]:
            cell.paragraphs[0].clear()
            set_paragraph_format(cell.paragraphs[0], alignment=align, space_after=0, line_spacing=1.15)
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=11, bold=(align == WD_ALIGN_PARAGRAPH.CENTER))
            set_cell_borders(cell)

    # Largeurs approximatives
    for row in table.rows:
        row.cells[0].width = Cm(3.5)
        row.cells[1].width = Cm(12.5)

    caption = doc.add_paragraph()
    set_paragraph_format(caption, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=6)
    run = caption.add_run("Tableau 1 — Liste des sigles et abréviations")
    set_run_font(run, size=10, italic=True)


def add_toc_entry(doc, title, level=0):
    """Entrée manuelle du sommaire détaillé (avec points de suite)."""
    p = doc.add_paragraph()
    indent = 0.5 * level
    set_paragraph_format(p, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=2, line_spacing=1.15)
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )

    size = 12 if level == 0 else 11
    bold = level == 0
    run = p.add_run(title)
    set_run_font(run, size=size, bold=bold)
    # Pas de numéro de page encore (corps non fourni) — points de suite uniquement
    run2 = p.add_run("\t")
    set_run_font(run2, size=size)
    return p


def build_listes_tableaux_figures(doc):
    add_page_break(doc)
    add_section_title(doc, "Liste des tableaux")

    p = doc.add_paragraph()
    set_paragraph_format(p, space_after=8, first_line_indent=0)
    run = p.add_run(
        "La liste des tableaux sera complétée au fur et à mesure de l’intégration des chapitres "
        "(légendes professionnelles sous chaque tableau)."
    )
    set_run_font(run, size=12)

    # Première entrée déjà créée
    entry = doc.add_paragraph()
    set_paragraph_format(entry, space_after=4, line_spacing=1.15)
    entry.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = entry.add_run("Tableau 1 — Liste des sigles et abréviations")
    set_run_font(run, size=11)
    run2 = entry.add_run("\t")
    set_run_font(run2, size=11)

    note = doc.add_paragraph()
    set_paragraph_format(note, space_before=8, space_after=6)
    run = note.add_run(
        "[Emplacement : les tableaux suivants (budget, risques, RACI, KPI, etc.) seront ajoutés "
        "lors de l’intégration des chapitres correspondants.]"
    )
    set_run_font(run, size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))

    add_page_break(doc)
    add_section_title(doc, "Liste des figures")

    p = doc.add_paragraph()
    set_paragraph_format(p, space_after=8, first_line_indent=0)
    run = p.add_run(
        "La liste des figures sera complétée au fur et à mesure de l’intégration des chapitres "
        "(légendes professionnelles sous chaque figure)."
    )
    set_run_font(run, size=12)

    note = doc.add_paragraph()
    set_paragraph_format(note, space_before=4, space_after=6)
    run = note.add_run("Information à compléter")
    set_run_font(run, size=12, italic=True)

    note2 = doc.add_paragraph()
    set_paragraph_format(note2, space_before=8, space_after=6)
    run = note2.add_run(
        "[Emplacements suggérés pour les figures à venir : logo ReadyToGo, carte des zones pilotes "
        "à Tanger, schéma du concept de service, WBS, roadmap 2026-2030, diagramme de Gantt, "
        "réseau PERT, matrice pouvoir/intérêt, matrice de criticité des risques, tableau de bord KPI.]"
    )
    set_run_font(run, size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))


def build_sommaire(doc):
    add_page_break(doc)
    add_section_title(doc, "Sommaire")

    note = doc.add_paragraph()
    set_paragraph_format(note, space_after=10, first_line_indent=0)
    run = note.add_run(
        "Sommaire détaillé du rapport. Les numéros de page seront actualisés automatiquement "
        "dans Microsoft Word (clic droit sur le sommaire > Mettre à jour les champs) une fois "
        "l’ensemble des chapitres intégré."
    )
    set_run_font(run, size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))

    # Champ TOC automatique
    toc_para = doc.add_paragraph()
    set_paragraph_format(toc_para, space_after=12)
    add_toc_field(toc_para)

    sep = doc.add_paragraph()
    set_paragraph_format(sep, space_before=6, space_after=10)
    run = sep.add_run("Structure détaillée proposée")
    set_run_font(run, size=12, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    structure = [
        ("Introduction générale", 0),
        ("Chapitre 1 : Présentation générale du projet ReadyToGo", 0),
        ("1.1. Origine de l’idée", 1),
        ("1.2. Présentation du concept", 1),
        ("1.3. Problème de mobilité à résoudre", 1),
        ("1.4. Objectifs du projet", 1),
        ("1.5. Valeur ajoutée du service", 1),
        ("1.6. Identité de marque ReadyToGo", 1),
        ("Chapitre 2 : Cadrage du projet", 0),
        ("2.1. Note de cadrage", 1),
        ("2.2. Problématique", 1),
        ("2.3. Objectifs SMART", 1),
        ("2.4. Périmètre du projet", 1),
        ("2.5. Hors périmètre", 1),
        ("2.6. Contraintes principales", 1),
        ("2.7. Livrables attendus", 1),
        ("Chapitre 3 : Étude du contexte et faisabilité", 0),
        ("3.1. Mobilité urbaine à Tanger", 1),
        ("3.2. Enjeux liés à la Coupe du Monde 2030", 1),
        ("3.3. Analyse des alternatives de transport", 1),
        ("3.4. Faisabilité technique", 1),
        ("3.5. Faisabilité économique", 1),
        ("3.6. Faisabilité juridique et réglementaire", 1),
        ("Chapitre 4 : Parties prenantes et gouvernance", 0),
        ("4.1. Identification des parties prenantes", 1),
        ("4.2. Analyse des attentes", 1),
        ("4.3. Matrice pouvoir / intérêt", 1),
        ("4.4. Matrice RACI", 1),
        ("4.5. Organisation de l’équipe projet", 1),
        ("Chapitre 5 : Méthodologie de management du projet", 0),
        ("5.1. Comparaison classique / agile", 1),
        ("5.2. Justification du choix hybride", 1),
        ("5.3. Phase classique du projet", 1),
        ("5.4. Phase agile du projet", 1),
        ("5.5. Organisation des sprints", 1),
        ("Chapitre 6 : Planification du projet", 0),
        ("6.1. Work Breakdown Structure", 1),
        ("6.2. Roadmap 2026-2030", 1),
        ("6.3. Jalons et livrables", 1),
        ("6.4. Diagramme de Gantt", 1),
        ("6.5. Réseau PERT", 1),
        ("6.6. Chemin critique", 1),
        ("Chapitre 7 : Budget et ressources", 0),
        ("7.1. Hypothèses budgétaires", 1),
        ("7.2. Coût des trottinettes", 1),
        ("7.3. Coût des bornes", 1),
        ("7.4. Coût de l’application mobile", 1),
        ("7.5. Ressources humaines", 1),
        ("7.6. Budget global prévisionnel", 1),
        ("7.7. Plan de charge", 1),
        ("Chapitre 8 : Gestion des risques", 0),
        ("8.1. Identification des risques", 1),
        ("8.2. Analyse probabilité / impact", 1),
        ("8.3. Registre des risques", 1),
        ("8.4. Matrice de criticité", 1),
        ("8.5. Plans d’action", 1),
        ("8.6. Suivi des risques", 1),
        ("Chapitre 9 : Suivi, contrôle et indicateurs KPI", 0),
        ("9.1. Objectifs du suivi", 1),
        ("9.2. KPI techniques", 1),
        ("9.3. KPI financiers", 1),
        ("9.4. KPI opérationnels", 1),
        ("9.5. KPI satisfaction utilisateur", 1),
        ("9.6. Tableau de bord projet", 1),
        ("Chapitre 10 : Communication et conduite du changement", 0),
        ("10.1. Plan de communication", 1),
        ("10.2. Communication interne", 1),
        ("10.3. Communication avec les partenaires", 1),
        ("10.4. Sensibilisation des utilisateurs", 1),
        ("10.5. Adoption du service", 1),
        ("Chapitre 11 : Scénarios réalistes d’utilisation", 0),
        ("11.1. Scénario étudiant", 1),
        ("11.2. Scénario habitant", 1),
        ("11.3. Scénario touriste", 1),
        ("11.4. Scénario jour de match", 1),
        ("11.5. Scénario maintenance", 1),
        ("Chapitre 12 : Livraison, clôture et perspectives", 0),
        ("12.1. Livrables finaux", 1),
        ("12.2. Recette du projet", 1),
        ("12.3. Bilan prévisionnel", 1),
        ("12.4. Limites du projet", 1),
        ("12.5. Perspectives d’évolution", 1),
        ("Conclusion générale", 0),
        ("Bibliographie", 0),
        ("Annexes", 0),
    ]

    for title, level in structure:
        add_toc_entry(doc, title, level)


def add_body_placeholder(doc):
    """Page de transition avant le corps du rapport (à compléter page par page)."""
    add_page_break(doc)
    heading = doc.add_heading("Introduction générale", level=1)
    for run in heading.runs:
        set_run_font(run, size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line_indent=0)
    run = p.add_run("Information à compléter")
    set_run_font(run, size=12, italic=True)

    note = doc.add_paragraph()
    set_paragraph_format(note, space_before=12)
    run = note.add_run(
        "[Le corps du rapport (introduction et chapitres) sera intégré progressivement "
        "au fur et à mesure de la réception du contenu validé, page par page.]"
    )
    set_run_font(run, size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))


def main():
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    setup_section(section, different_first_page=True)

    build_cover_page(doc)
    build_remerciements(doc)
    build_resume(doc)
    build_abstract(doc)
    build_sigles(doc)
    build_listes_tableaux_figures(doc)
    build_sommaire(doc)
    add_body_placeholder(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))

    # Copie artifact téléchargeable
    artifact = Path("/opt/cursor/artifacts") / OUTPUT.name
    artifact.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(artifact))

    print(f"Document créé : {OUTPUT}")
    print(f"Artifact : {artifact}")


if __name__ == "__main__":
    main()
