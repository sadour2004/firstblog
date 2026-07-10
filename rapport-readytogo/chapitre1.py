"""Chapitre 1 — Présentation générale du projet ReadyToGo."""

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


def _import_helpers():
    from generer_pages_preliminaires import (
        add_body_paragraph,
        add_page_break,
        set_cell_borders,
        set_cell_shading,
        set_paragraph_format,
        set_run_font,
    )
    return {
        "add_body_paragraph": add_body_paragraph,
        "add_page_break": add_page_break,
        "set_cell_borders": set_cell_borders,
        "set_cell_shading": set_cell_shading,
        "set_paragraph_format": set_paragraph_format,
        "set_run_font": set_run_font,
    }


def add_rich_paragraph(doc, text, first_indent=0.75, helpers=None):
    """Paragraphe justifié avec segments **gras**."""
    h = helpers
    p = doc.add_paragraph()
    h["set_paragraph_format"](p, first_line_indent=first_indent, space_after=10)
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        h["set_run_font"](run, size=12, bold=(i % 2 == 1))
    return p


def add_heading2(doc, text, helpers):
    heading = doc.add_heading(text, level=2)
    for run in heading.runs:
        helpers["set_run_font"](run, size=14, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    helpers["set_paragraph_format"](
        heading, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=8
    )
    return heading


def add_heading3(doc, text, helpers):
    heading = doc.add_heading(text, level=3)
    for run in heading.runs:
        helpers["set_run_font"](run, size=12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
    helpers["set_paragraph_format"](
        heading, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=6
    )
    return heading


def add_quote(doc, text, helpers):
    p = doc.add_paragraph()
    helpers["set_paragraph_format"](
        p, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=8, space_after=10, first_line_indent=0
    )
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.right_indent = Cm(1.0)
    run = p.add_run(text)
    helpers["set_run_font"](run, size=12, italic=True)
    return p


def add_caption(doc, text, helpers):
    p = doc.add_paragraph()
    helpers["set_paragraph_format"](
        p, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=12, line_spacing=1.15
    )
    run = p.add_run(text)
    helpers["set_run_font"](run, size=10, italic=True)
    return p


def add_figure_placeholder(doc, title, description, helpers):
    box = doc.add_paragraph()
    helpers["set_paragraph_format"](
        box, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=4, line_spacing=1.15
    )
    run = box.add_run(f"[Emplacement — {title}]")
    helpers["set_run_font"](run, size=10, italic=True, color=RGBColor(0x66, 0x66, 0x66))

    desc = doc.add_paragraph()
    helpers["set_paragraph_format"](
        desc, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4, line_spacing=1.15
    )
    run = desc.add_run(description)
    helpers["set_run_font"](run, size=9, italic=True, color=RGBColor(0x88, 0x88, 0x88))


def add_academic_table(doc, headers, rows, helpers, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, label in enumerate(headers):
        hdr_cells[i].paragraphs[0].clear()
        helpers["set_paragraph_format"](
            hdr_cells[i].paragraphs[0],
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            space_after=0,
            line_spacing=1.15,
        )
        run = hdr_cells[i].paragraphs[0].add_run(label)
        helpers["set_run_font"](run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        helpers["set_cell_shading"](hdr_cells[i], "1F4E79")
        helpers["set_cell_borders"](hdr_cells[i])

    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].paragraphs[0].clear()
            helpers["set_paragraph_format"](
                cells[i].paragraphs[0],
                alignment=WD_ALIGN_PARAGRAPH.LEFT,
                space_after=0,
                line_spacing=1.15,
            )
            run = cells[i].paragraphs[0].add_run(str(value))
            helpers["set_run_font"](run, size=10, bold=(i == 0))
            helpers["set_cell_borders"](cells[i])

    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)

    return table


def build_chapitre_1(doc):
    h = _import_helpers()
    h["add_page_break"](doc)

    title = doc.add_heading("Chapitre 1 : Présentation générale du projet ReadyToGo", level=1)
    for run in title.runs:
        h["set_run_font"](run, size=16, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    h["set_paragraph_format"](title, alignment=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=12)

    # --- 1.1 ---
    add_heading2(doc, "1.1. Origine de l’idée du projet", h)

    add_rich_paragraph(
        doc,
        "Le projet **ReadyToGo** est né d’un constat lié à la mobilité urbaine à Tanger. "
        "La ville connaît une croissance importante sur les plans économique, touristique et "
        "démographique. Cette évolution entraîne une augmentation des déplacements quotidiens, "
        "une pression sur les axes routiers et des difficultés de circulation, notamment dans "
        "les zones à forte fréquentation comme le centre-ville, la corniche, la gare Tanger-Ville, "
        "les quartiers touristiques et les axes menant vers le Grand Stade de Tanger.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "La fiche initiale du projet met en évidence un problème central : certains trajets "
        "urbains peuvent devenir longs et peu efficaces en voiture, surtout en période de "
        "congestion. Elle indique notamment que le trajet entre le centre-ville et le Grand Stade "
        "peut prendre un temps important en voiture, ce qui montre la nécessité de proposer des "
        "solutions complémentaires de mobilité.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Dans cette perspective, l’idée de ReadyToGo consiste à proposer un service de "
        "**trottinettes électriques en libre-service**, permettant aux habitants, étudiants, "
        "travailleurs et visiteurs de réaliser des trajets courts ou moyens de manière plus "
        "rapide, flexible et accessible. Le projet s’inscrit également dans une vision plus "
        "large liée à la préparation de Tanger aux grands événements internationaux, notamment "
        "la Coupe du Monde 2030.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "ReadyToGo ne se présente donc pas comme un simple service de location de trottinettes. "
        "Il s’agit d’un projet de mobilité urbaine qui cherche à améliorer l’expérience de "
        "déplacement en ville, réduire la dépendance aux véhicules individuels sur les courtes "
        "distances et compléter les solutions de transport existantes.",
        helpers=h,
    )

    # --- 1.2 ---
    add_heading2(doc, "1.2. Présentation du concept ReadyToGo", h)

    add_rich_paragraph(
        doc,
        "**ReadyToGo** est une application mobile associée à une flotte de trottinettes "
        "électriques disponibles dans plusieurs zones stratégiques de Tanger. L’utilisateur "
        "peut localiser une trottinette proche de lui, la déverrouiller à l’aide d’un QR code, "
        "effectuer son trajet, puis la stationner dans une zone autorisée.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Le service repose sur une logique simple : rendre les déplacements urbains plus "
        "rapides, plus pratiques et plus adaptés aux besoins actuels des usagers. La fiche de "
        "projet initiale présente déjà ce fonctionnement quotidien : téléchargement de "
        "l’application, scan du QR code, paiement numérique, utilisation de la trottinette, "
        "stationnement dans une zone autorisée, puis recharge et maintenance par une équipe "
        "locale.",
        helpers=h,
    )

    add_heading3(doc, "Tableau 1 — Synthèse du concept ReadyToGo", h)
    add_academic_table(
        doc,
        ["Élément", "Description"],
        [
            ("Nom du projet", "ReadyToGo"),
            ("Nature du projet", "Service de trottinettes électriques en libre-service"),
            ("Ville cible", "Tanger"),
            ("Public cible", "Habitants, étudiants, travailleurs, touristes, visiteurs du Mondial"),
            ("Support principal", "Application mobile"),
            ("Moyen de déplacement", "Trottinettes électriques géolocalisées"),
            ("Mode de paiement", "Carte bancaire, portefeuille mobile ou pass"),
            ("Objectif principal", "Améliorer la mobilité urbaine courte distance"),
            ("Positionnement", "Solution complémentaire au bus, taxi, marche et voiture"),
        ],
        h,
        col_widths=[5.0, 11.0],
    )
    add_caption(doc, "Tableau 1 — Synthèse du concept ReadyToGo", h)

    # --- 1.3 ---
    add_heading2(doc, "1.3. Problème à résoudre", h)

    add_rich_paragraph(
        doc,
        "La problématique principale du projet peut être formulée ainsi :",
        first_indent=0,
        helpers=h,
    )
    add_quote(
        doc,
        "Comment proposer à Tanger une solution de mobilité urbaine flexible, rapide et "
        "accessible, capable de réduire les difficultés de déplacement sur les trajets courts "
        "et moyens, tout en accompagnant la ville dans sa préparation aux grands événements "
        "internationaux ?",
        h,
    )
    add_rich_paragraph(
        doc,
        "Cette problématique est liée à plusieurs difficultés concrètes.",
        helpers=h,
    )

    add_heading3(doc, "Tableau 2 — Problèmes identifiés et réponses proposées par ReadyToGo", h)
    add_academic_table(
        doc,
        ["Problème identifié", "Conséquence", "Réponse proposée par ReadyToGo"],
        [
            (
                "Embouteillages dans certaines zones urbaines",
                "Perte de temps, stress, baisse de fluidité",
                "Proposer une alternative légère pour les courts trajets",
            ),
            (
                "Difficulté de stationnement",
                "Retards, coûts supplémentaires",
                "Utilisation de zones de stationnement dédiées",
            ),
            (
                "Taxis parfois coûteux ou indisponibles",
                "Mobilité moins accessible",
                "Tarification intermédiaire entre bus et taxi",
            ),
            (
                "Bus moins flexible sur certains trajets",
                "Temps d’attente et itinéraires fixes",
                "Disponibilité à la demande via application",
            ),
            (
                "Hausse attendue des visiteurs en 2030",
                "Pression sur les transports",
                "Offre adaptée aux touristes avec pass journée/semaine",
            ),
            (
                "Besoin de mobilité écologique",
                "Pollution et congestion",
                "Déplacement électrique à faible émission locale",
            ),
        ],
        h,
        col_widths=[5.2, 5.0, 5.8],
    )
    add_caption(doc, "Tableau 2 — Problèmes identifiés et réponses proposées par ReadyToGo", h)

    add_rich_paragraph(
        doc,
        "Le projet répond donc à une double logique : améliorer la mobilité quotidienne des "
        "habitants et renforcer l’attractivité de Tanger comme ville moderne et intelligente.",
        helpers=h,
    )

    # --- 1.4 ---
    add_heading2(doc, "1.4. Objectifs généraux du projet", h)

    add_rich_paragraph(
        doc,
        "ReadyToGo poursuit plusieurs objectifs complémentaires.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Le premier objectif est **opérationnel** : mettre à disposition des utilisateurs un "
        "service simple et fiable de trottinettes électriques.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Le deuxième objectif est **urbain** : contribuer à la réduction de la congestion dans "
        "certaines zones à forte circulation.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Le troisième objectif est **économique** : créer un modèle de service viable, basé sur "
        "les paiements à l’usage, les abonnements, les pass touristiques et les partenariats.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Le quatrième objectif est **stratégique** : accompagner la transformation de Tanger "
        "dans la perspective de la Coupe du Monde 2030.",
        helpers=h,
    )

    add_heading3(doc, "Tableau 3 — Objectifs généraux de ReadyToGo", h)
    add_academic_table(
        doc,
        ["Type d’objectif", "Description"],
        [
            ("Objectif de mobilité", "Faciliter les déplacements courts et moyens à Tanger"),
            ("Objectif utilisateur", "Offrir un service simple, rapide, accessible et flexible"),
            ("Objectif environnemental", "Encourager une mobilité électrique et moins polluante"),
            ("Objectif économique", "Construire un modèle rentable et durable"),
            ("Objectif territorial", "Participer à l’image moderne et connectée de Tanger"),
            ("Objectif événementiel", "Préparer une solution adaptée aux flux touristiques de 2030"),
        ],
        h,
        col_widths=[5.0, 11.0],
    )
    add_caption(doc, "Tableau 3 — Objectifs généraux de ReadyToGo", h)

    # --- 1.5 ---
    add_heading2(doc, "1.5. Valeur ajoutée du service", h)

    add_rich_paragraph(
        doc,
        "La valeur ajoutée de ReadyToGo repose sur sa capacité à combiner **rapidité**, "
        "**flexibilité**, **accessibilité** et **simplicité d’usage**.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Contrairement au bus, qui dépend d’horaires et de lignes fixes, ReadyToGo permet à "
        "l’utilisateur de choisir son point de départ et son point d’arrivée dans une zone "
        "autorisée. Contrairement au taxi, le coût peut rester plus prévisible et plus adapté "
        "aux petits trajets. Contrairement à la voiture personnelle, la trottinette ne "
        "nécessite pas de stationnement classique et peut être plus efficace dans les zones "
        "encombrées.",
        helpers=h,
    )

    add_heading3(doc, "Tableau 4 — Comparaison entre ReadyToGo et les alternatives de transport", h)
    add_academic_table(
        doc,
        ["Critère", "Bus", "Taxi", "Voiture personnelle", "Marche", "ReadyToGo"],
        [
            ("Coût", "Faible", "Moyen", "Élevé", "Gratuit", "Moyen/faible"),
            ("Rapidité sur court trajet", "Moyenne", "Variable", "Variable", "Faible", "Élevée"),
            ("Flexibilité", "Faible", "Bonne", "Bonne", "Bonne", "Très bonne"),
            ("Problème de stationnement", "Non", "Non", "Oui", "Non", "Limité"),
            ("Adapté aux touristes", "Moyen", "Oui", "Non", "Moyen", "Oui"),
            ("Impact écologique local", "Moyen", "Moyen", "Élevé", "Nul", "Faible"),
        ],
        h,
        col_widths=[3.8, 2.2, 2.2, 2.8, 2.2, 2.8],
    )
    add_caption(doc, "Tableau 4 — Comparaison entre ReadyToGo et les alternatives de transport", h)

    add_rich_paragraph(
        doc,
        "La fiche initiale positionne déjà le service entre le bus et le petit taxi : plus cher "
        "que le bus, mais plus rapide et plus flexible ; moins cher qu’un taxi pour certains "
        "trajets courts.",
        helpers=h,
    )

    # --- 1.6 ---
    add_heading2(doc, "1.6. Identité de marque ReadyToGo", h)

    add_rich_paragraph(
        doc,
        "Le nom **ReadyToGo** signifie littéralement “prêt à partir”. Il exprime l’idée d’un "
        "service immédiat, rapide et simple. L’utilisateur ouvre l’application, trouve une "
        "trottinette, scanne le QR code et commence son trajet.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "Ce nom présente plusieurs avantages :",
        first_indent=0,
        helpers=h,
    )

    add_academic_table(
        doc,
        ["Élément", "Intérêt"],
        [
            ("Court et mémorisable", "Facile à retenir par les habitants et touristes"),
            ("Compréhensible en anglais", "Adapté au contexte international de 2030"),
            ("Dynamique", "Donne une idée de mouvement et de rapidité"),
            ("Moderne", "Compatible avec une application mobile et une marque digitale"),
            ("Évolutif", "Peut être utilisé plus tard dans d’autres villes marocaines"),
        ],
        h,
        col_widths=[5.0, 11.0],
    )
    add_caption(doc, "Tableau — Avantages du nom ReadyToGo", h)

    add_heading3(doc, "Slogan proposé", h)
    add_quote(doc, "ReadyToGo — Bougez plus vite, vivez Tanger autrement.", h)

    add_rich_paragraph(doc, "Autres slogans possibles :", first_indent=0, helpers=h)
    add_academic_table(
        doc,
        ["Slogan", "Positionnement"],
        [
            ("ReadyToGo, Tanger en mouvement", "Image urbaine et dynamique"),
            ("ReadyToGo, plus rapide que le trafic", "Accent sur la rapidité"),
            ("ReadyToGo, votre trajet commence ici", "Accent sur l’expérience utilisateur"),
            ("ReadyToGo, la mobilité simple à Tanger", "Accent sur la simplicité"),
        ],
        h,
        col_widths=[8.0, 8.0],
    )
    add_caption(doc, "Tableau — Slogans possibles pour ReadyToGo", h)

    add_figure_placeholder(
        doc,
        "Figure 1 : Logo conceptuel ReadyToGo",
        "À réaliser avec Canva, puis à intégrer dans le rapport final.",
        h,
    )
    add_caption(doc, "Figure 1 — Logo conceptuel ReadyToGo (à intégrer)", h)

    # --- 1.7 ---
    add_heading2(doc, "1.7. Fonctionnement général du service", h)

    add_rich_paragraph(
        doc,
        "Le fonctionnement de ReadyToGo peut être résumé en sept étapes.",
        helpers=h,
    )

    add_heading3(doc, "Figure 2 — Parcours utilisateur ReadyToGo", h)
    add_figure_placeholder(
        doc,
        "Figure 2 : Parcours utilisateur ReadyToGo",
        "Télécharger l’application → Créer un compte → Localiser une trottinette → Scanner le QR code → "
        "Effectuer le trajet → Stationner dans une zone autorisée → Paiement et fin du trajet\n"
        "Outil conseillé : Miro, Lucidchart, Canva ou Figma.",
        h,
    )
    add_caption(doc, "Figure 2 — Parcours utilisateur ReadyToGo (à intégrer)", h)

    add_heading3(doc, "Tableau 5 — Parcours utilisateur détaillé", h)
    add_academic_table(
        doc,
        ["Étape", "Action utilisateur", "Fonction du système"],
        [
            ("1", "Télécharger l’application", "Accès au service"),
            ("2", "Créer un compte", "Identification et sécurité"),
            ("3", "Chercher une trottinette", "Carte avec disponibilité en temps réel"),
            ("4", "Scanner le QR code", "Déverrouillage de la trottinette"),
            ("5", "Réaliser le trajet", "Suivi GPS et calcul du coût"),
            ("6", "Stationner correctement", "Vérification de la zone autorisée"),
            ("7", "Terminer le trajet", "Paiement automatique et reçu numérique"),
        ],
        h,
        col_widths=[2.0, 6.0, 8.0],
    )
    add_caption(doc, "Tableau 5 — Parcours utilisateur détaillé", h)

    # --- 1.8 ---
    add_heading2(doc, "1.8. Fonctionnement opérationnel quotidien", h)

    add_rich_paragraph(
        doc,
        "Le projet ne repose pas uniquement sur une application mobile. Il nécessite aussi une "
        "organisation opérationnelle. Une équipe locale doit assurer la disponibilité, la "
        "recharge, la maintenance, le nettoyage et la redistribution des trottinettes.",
        helpers=h,
    )

    add_heading3(doc, "Tableau 6 — Organisation opérationnelle quotidienne", h)
    add_academic_table(
        doc,
        ["Activité", "Description", "Responsable"],
        [
            ("Recharge", "Recharger les batteries chaque nuit ou selon besoin", "Équipe opérationnelle"),
            ("Maintenance", "Réparer les trottinettes endommagées", "Technicien maintenance"),
            (
                "Redistribution",
                "Déplacer les trottinettes vers les zones demandées",
                "Agents terrain",
            ),
            ("Assistance utilisateur", "Répondre aux réclamations", "Support client"),
            (
                "Suivi de flotte",
                "Contrôler disponibilité, pannes et localisation",
                "Responsable opérations",
            ),
            (
                "Reporting",
                "Produire des indicateurs quotidiens",
                "Chef de projet / Data analyst",
            ),
        ],
        h,
        col_widths=[4.0, 7.5, 4.5],
    )
    add_caption(doc, "Tableau 6 — Organisation opérationnelle quotidienne", h)

    add_rich_paragraph(
        doc,
        "Cette dimension opérationnelle est essentielle, car elle montre que ReadyToGo n’est "
        "pas seulement un produit numérique, mais un **service complet** combinant technologie, "
        "logistique, maintenance, relation client et gestion urbaine.",
        helpers=h,
    )

    # --- 1.9 ---
    add_heading2(doc, "1.9. Justification du choix de Tanger", h)

    add_rich_paragraph(
        doc,
        "Tanger est une ville pertinente pour lancer ReadyToGo pour plusieurs raisons.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "D’abord, elle possède une forte dynamique économique et touristique. Ensuite, elle "
        "dispose de zones urbaines attractives où les courts trajets sont fréquents : corniche, "
        "centre-ville, gare, zones hôtelières, quartiers étudiants et espaces touristiques. "
        "Enfin, la perspective de la Coupe du Monde 2030 renforce l’intérêt de tester des "
        "solutions innovantes de mobilité.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "La fiche initiale du projet propose d’ailleurs un déploiement progressif à Tanger : "
        "test dans la corniche et le centre-ville, extension vers la gare, l’aéroport et les "
        "quartiers hôteliers, puis connexion avec le Grand Stade avant 2030.",
        helpers=h,
    )

    add_heading3(doc, "Tableau 7 — Zones potentielles de déploiement à Tanger", h)
    add_academic_table(
        doc,
        ["Zone", "Intérêt pour le projet"],
        [
            ("Corniche", "Forte fréquentation, tourisme, loisirs"),
            ("Centre-ville", "Mobilité quotidienne, commerces, services"),
            ("Gare Tanger-Ville", "Connexion avec le train et les visiteurs"),
            ("Quartiers hôteliers", "Besoin des touristes et voyageurs"),
            ("Aéroport Ibn Battouta", "Extension stratégique à moyen terme"),
            ("Grand Stade de Tanger", "Enjeu majeur pour les événements sportifs"),
            ("Zones universitaires", "Potentiel pour les étudiants"),
        ],
        h,
        col_widths=[5.0, 11.0],
    )
    add_caption(doc, "Tableau 7 — Zones potentielles de déploiement à Tanger", h)

    add_figure_placeholder(
        doc,
        "Figure 3 : Carte conceptuelle des zones pilotes à Tanger",
        "À réaliser avec Google My Maps ou Canva, puis à intégrer dans le rapport final.",
        h,
    )
    add_caption(doc, "Figure 3 — Carte conceptuelle des zones pilotes à Tanger (à intégrer)", h)

    add_figure_placeholder(
        doc,
        "Figure 4 : Comparaison visuelle des moyens de transport",
        "À réaliser avec Canva, puis à intégrer dans le rapport final.",
        h,
    )
    add_caption(doc, "Figure 4 — Comparaison visuelle des moyens de transport (à intégrer)", h)

    # --- 1.10 ---
    add_heading2(doc, "1.10. Conclusion du chapitre", h)

    add_rich_paragraph(
        doc,
        "Ce premier chapitre a permis de présenter le projet ReadyToGo dans sa globalité. Le "
        "projet répond à un besoin réel de mobilité urbaine à Tanger, en proposant une solution "
        "de trottinettes électriques en libre-service. Il se distingue par sa flexibilité, son "
        "accessibilité, son potentiel touristique et sa capacité à compléter les transports "
        "existants.",
        helpers=h,
    )
    add_rich_paragraph(
        doc,
        "ReadyToGo est également un projet pertinent sur le plan du management de projet, car "
        "il nécessite un cadrage clair, une planification progressive, une gestion des risques, "
        "un budget prévisionnel, une organisation opérationnelle et un suivi par indicateurs. "
        "Il s’inscrit donc pleinement dans les objectifs du cours, qui vise à construire un "
        "projet de bout en bout en respectant les contraintes de coût, de délai et de qualité.",
        helpers=h,
    )
